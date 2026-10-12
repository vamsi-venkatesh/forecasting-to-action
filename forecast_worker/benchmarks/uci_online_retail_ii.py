"""Public numerical benchmark: UCI Online Retail II → weekly net unit sales per product (brief §5).

Source: Chen, D. (2012). Online Retail II. UCI Machine Learning Repository. https://doi.org/10.24432/C5CG6D
Licence: CC BY 4.0. File: online_retail_ii.zip, 45,622,418 bytes, SHA-256
572e36277c2390fbfde10664750731e0a86f55e33470d91919085f0408e67bfb (downloaded 2026-09-28 from
https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip). The script refuses any other file.

What the data is and is not (recorded in the manifest, never filled in):
- Transactions of a UK online gift-ware retailer, mostly wholesale customers, 2009-12-01 to 2011-12-09.
- Quantities are units per invoice line; prices in GBP. Cancellations are invoices starting with "C"; they are
  counted separately and netted into the week of the cancellation (the data does not link most of them to
  the original line). Negative quantities on non-"C" invoices (stock adjustments) are excluded and counted.
- NOT in the data: stock levels, receipts, promotions, availability. So sales here are observed sales; stockouts
  cannot be detected, and nothing about inventory is inferred.
- The two yearly sheets overlap in early December 2010: identical lines in both are counted once.
- Non-product codes (postage, fees, manual adjustments, samples) are excluded and listed.
- available_at is an ASSUMPTION (week end + 1 day) for the lag-aware evaluation; the source has no such field.

    python -m forecast_worker.benchmarks.uci_online_retail_ii ZIP_OR_XLSX OUT_DIR [--top 50]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ZIP_SHA256 = "572e36277c2390fbfde10664750731e0a86f55e33470d91919085f0408e67bfb"
NON_PRODUCT = re.compile(r"^(POST|DOT|M|D|C2|CRUK|BANK CHARGES|AMAZONFEE|ADJUST\d*|B|S|PADS|TEST\d*|GIFT_.*|SP\d+)$", re.I)
BENCHMARK_VERSION = "uci-online-retail-ii-weekly-2"   # v2: series chosen on training-period data only; no
                                                      # zeros before a product's first sale; partial first week dropped
SELECTION_OFFSET_WEEKS = 24   # the default protocol's first evaluation origin: holdout 8 + 12 origins + horizon 4 + gap 1 - 1


def _monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def ingest(src: Path, out: Path, top: int = 50, selection_offset: int = SELECTION_OFFSET_WEEKS) -> dict:
    raw = src.read_bytes()
    if src.suffix == ".zip":
        if hashlib.sha256(raw).hexdigest() != ZIP_SHA256:
            raise SystemExit("not the recorded UCI Online Retail II file (SHA-256 differs)")
        with zipfile.ZipFile(src) as z:
            name = next(n for n in z.namelist() if n.lower().endswith(".xlsx"))
            xlsx = out / "_source.xlsx"
            out.mkdir(parents=True, exist_ok=True)
            xlsx.write_bytes(z.read(name))
    else:
        xlsx = src
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)   # cached values only: no formula or macro runs
    seen: set[tuple] = set()
    counts = Counter()
    weekly: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    cancelled: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    excluded_codes = Counter()
    first = last = None
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        ix = {h: i for i, h in enumerate(header)}
        for r in rows:
            counts["rows"] += 1
            key = tuple(r)
            if key in seen:
                counts["duplicate_lines_across_or_within_sheets"] += 1
                continue
            seen.add(key)
            inv, code, qty, when = str(r[ix["Invoice"]]), str(r[ix["StockCode"]]).strip(), r[ix["Quantity"]], r[ix["InvoiceDate"]]
            if not isinstance(when, datetime) or not isinstance(qty, (int, float)):
                counts["unparsed"] += 1
                continue
            first, last = min(first or when, when), max(last or when, when)
            if NON_PRODUCT.match(code):
                excluded_codes[code] += 1
                continue
            wk = _monday(when.date())
            if inv.upper().startswith("C"):
                counts["cancellation_lines"] += 1
                cancelled[code][wk] += -abs(qty)
            elif qty < 0:
                counts["negative_non_cancellation_lines_excluded"] += 1
            else:
                counts["sale_lines"] += 1
                weekly[code][wk] += qty
    full_weeks = [d for d in _weeks(_monday(first.date()), _monday(last.date()))]
    last_complete = full_weeks[-2]                        # the final week is partial (data ends on a Friday)
    first_complete = full_weeks[0] if first.date() == full_weeks[0] else full_weeks[1]   # data starts on a Tuesday
    full_weeks = [w for w in full_weeks if first_complete <= w <= last_complete]
    # selection uses ONLY weeks up to the first evaluation origin: never later sales, never survival into the test
    select_until = last_complete - timedelta(weeks=selection_offset)
    totals = sorted(((sum(q for w, q in v.items() if first_complete <= w <= select_until), c)
                     for c, v in weekly.items()), reverse=True)
    chosen = [c for t, c in totals[:top] if t > 0]
    first_sale = {c: min(w for w, q in weekly[c].items() if q > 0 and w >= first_complete) for c in chosen}
    out.mkdir(parents=True, exist_ok=True)
    path = out / "sales_uci_weekly.csv"
    with path.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["source_record_id", "series_id", "product_id", "region_id", "week_start", "available_at",
                    "unit", "returned", "source_revision", "quantity"])
        for code in chosen:
            for wk in full_weeks:
                if wk < first_sale[code]:
                    continue                              # not yet sold: absent, never a zero
                net = weekly[code].get(wk, 0.0) + cancelled[code].get(wk, 0.0)
                avail = datetime.combine(wk + timedelta(days=7), datetime.min.time()).isoformat() + "+00:00"
                w.writerow([f"UCI-{code}-{wk.isoformat()}", code, code, "ALL", wk.isoformat(), avail, "units",
                            int(-cancelled[code].get(wk, 0.0)), 1, int(net)])
    manifest = {"benchmark": BENCHMARK_VERSION, "source": "UCI Online Retail II (Chen 2012), doi:10.24432/C5CG6D",
                "licence": "CC BY 4.0", "zip_sha256": ZIP_SHA256, "first_transaction": first.isoformat(),
                "last_transaction": last.isoformat(), "last_complete_week": last_complete.isoformat(),
                "weeks": len(full_weeks), "series": len(chosen),
                "selection": f"top {top} product codes by gross units sold (all countries) in the weeks "
                             f"{first_complete.isoformat()} .. {select_until.isoformat()} only (before the first "
                             "evaluation origin): later sales and survival into the test period play no part",
                "target": "net units per week = units sold minus units on cancellation invoices, netted into the week "
                          "of the cancellation (the source rarely links a cancellation to its sale)",
                "first_complete_week": first_complete.isoformat(), "selection_until": select_until.isoformat(),
                "series_start": "each series starts at its first week with sales; earlier weeks are absent, not zero",
                "currency": "GBP (retail selling prices; not used here and never relabelled EUR)",
                "counts": dict(counts), "excluded_non_product_codes": dict(excluded_codes.most_common(20)),
                "not_in_data": ["stock levels", "receipts", "promotions", "availability", "stockouts"],
                "assumptions": ["exact duplicate lines (all eight fields equal) are counted once: the two sheets "
                                "overlap in December 2010 and the source also repeats some lines within a sheet; "
                                "a genuinely repeated identical purchase would be undercounted by this choice",
                                "available_at = week start + 7 days (the source has no availability time)",
                                "cancellations netted into the week they occur", "weeks with no line = 0 units "
                                "(a transaction log records every sale, so an empty week is no sales, not missing)"],
                "output_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1) + "\n")
    if src.suffix == ".zip":
        (out / "_source.xlsx").unlink()
    return manifest


def _weeks(a: date, b: date):
    d = a
    while d <= b:
        yield d
        d += timedelta(weeks=1)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--top", type=int, default=50)
    ap.add_argument("--selection-offset", type=int, default=SELECTION_OFFSET_WEEKS)
    a = ap.parse_args()
    m = ingest(Path(a.src), Path(a.out), a.top, a.selection_offset)
    print(json.dumps({k: m[k] for k in ("benchmark", "weeks", "series", "counts", "output_sha256")}, indent=1))

